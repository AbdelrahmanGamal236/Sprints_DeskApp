"""
Server protocol tests: authentication, server-side authorization, identity binding, error handling,
and export streaming. Each scenario from the Phase 1 audit (AUDIT_REPORT.md, Section 3) is replayed
and must now be rejected.
"""
import asyncio
import json
import sqlite3
import unittest
import uuid

from app.core.config import ROLE_HEAD, ROLE_ADMIN, ROLE_COO, ROLE_SUPER_ADMIN
from app.network import protocol as P
from app.network.server import ACTIONS
from tests.helpers import start_server, RawClient, raw_login, PASSWORD

H, A, C, S = ROLE_HEAD, ROLE_ADMIN, ROLE_COO, ROLE_SUPER_ADMIN
ALL = {H, A, C, S}

# The role table from the spec. Changing a permission must be a deliberate edit here.
EXPECTED_ROLES = {
    "GET_SNAPSHOT": ALL, "LOGOUT": ALL, "CHANGE_PASSWORD": ALL,
    "GET_DISCIPLINES": ALL, "GET_PROJECTS": ALL, "GET_TASKS": ALL, "GET_PROJECT_TIMELINE": ALL,
    "SUBMIT_LOG": {H}, "CALCULATE_NEXT_REVISION": {H}, "PROPOSE_TASK": {H},
    "REQUEST_TASK_DATE_CHANGE": {H}, "RESPOND_NOTIFICATION": {H},
    "GET_TASK_PROPOSALS": {H, A, S}, "GET_DATE_CHANGE_REQUESTS": {H, A, S},
    "GET_NOTIFICATIONS": {H, C, S},
    "GENERATE_NEXT_TASK_ID": {A, S}, "APPROVE_TASK_PROPOSAL": {A, S}, "REJECT_TASK_PROPOSAL": {A, S},
    "APPROVE_DATE_CHANGE": {A, S}, "REJECT_DATE_CHANGE": {A, S},
    "ADD_DISCIPLINE": {A, S}, "TOGGLE_HIDE_DISCIPLINE": {A, S},
    "CREATE_USER": {A, S}, "UPDATE_USER": {A, S}, "RESET_PASSWORD": {A, S},
    "ADD_PROJECT": {A, S}, "TOGGLE_HIDE_PROJECT": {A, S}, "DELETE_PROJECT": {A, S}, "DELETE_TASK": {A, S},
    "ADD_TASK": {A, S}, "UPDATE_TASK": {A, S}, "UPDATE_TASK_DATES": {A, S}, "TOGGLE_HIDE_TASK": {A, S},
    "GET_ANALYTICS": {A, S},
    "GET_RECENT_LOGS": {A, C, S}, "GET_USERS": {A, C, S}, "EXPORT_EXCEL_REQUEST": {A, C, S},
    "COO_NOTIFY": {C, S},
    "GET_SETTINGS": {S}, "UPDATE_SETTINGS": {S},
}
PUBLIC = {"PING", "LOGIN", "RESUME"}
ROLE_USER = {S: "sa", A: "adm", C: "coo1", H: "h_urb"}


def run(coro):
    return asyncio.run(coro)


def log_payload(task_id="T-URB", project_id="P1", **extra):
    p = {"log_id": f"LOG-{uuid.uuid4().hex.upper()}", "project_id": project_id, "task_id": task_id,
         "start_date": "2026-10-01", "end_date": "2026-10-01", "progress_stage": "In Progress"}
    p.update(extra)
    return p


class TestPermissionTable(unittest.TestCase):
    def test_registry_matches_spec_table(self):
        registered = {name: (set(spec.roles) if spec.roles else ALL) for name, spec in ACTIONS.items() if not spec.public}
        self.assertEqual(registered, EXPECTED_ROLES)
        self.assertEqual({n for n, s in ACTIONS.items() if s.public}, PUBLIC)

    def test_live_matrix_every_action_every_role(self):
        async def scenario():
            server = await start_server()
            try:
                anon = await RawClient(server.port).connect()
                clients = {role: await raw_login(server.port, uid) for role, uid in ROLE_USER.items()}
                for action_name, allowed in EXPECTED_ROLES.items():
                    if action_name in ("LOGOUT", "CHANGE_PASSWORD"):
                        continue
                    r = await anon.req(action_name, {})
                    self.assertEqual(r["code"], P.ERR_AUTH_REQUIRED, action_name)
                    for role, client in clients.items():
                        r = await client.req(action_name, {})
                        if role in allowed:
                            self.assertNotIn(r.get("code"), (P.ERR_FORBIDDEN, P.ERR_AUTH_REQUIRED),
                                             f"{role} should be allowed {action_name}: {r}")
                        else:
                            self.assertEqual(r.get("code"), P.ERR_FORBIDDEN, f"{role} must be denied {action_name}: {r}")
            finally:
                await server.stop()
        run(scenario())


class TestAuthentication(unittest.TestCase):
    def test_audit_probes_are_rejected(self):
        async def scenario():
            server = await start_server()
            try:
                anon = await RawClient(server.port).connect()
                # Probe 2: IDENTIFY no longer exists
                r = await anon.req("IDENTIFY", {"user_id": "x", "role": "SUPER_ADMIN"})
                self.assertEqual(r["code"], P.ERR_UNKNOWN_ACTION)
                # Probe 3: anonymous CREATE_USER
                r = await anon.req("CREATE_USER", {"user_id": "evil", "full_name": "x", "role": "COO"})
                self.assertEqual(r["code"], P.ERR_AUTH_REQUIRED)
                # Probe 4: anonymous forced password change
                r = await anon.req("CHANGE_PASSWORD", {"user_id": "sa", "new_password": "pwned1234", "force": True})
                self.assertEqual(r["code"], P.ERR_AUTH_REQUIRED)
                r = await anon.login("sa", "pwned1234")
                self.assertFalse(r["ok"])
                # Probe 5: anonymous data access
                r = await anon.req("GET_USERS", {})
                self.assertEqual(r["code"], P.ERR_AUTH_REQUIRED)
                self.assertIsNone(server.db.get_user("evil"))
            finally:
                await server.stop()
        run(scenario())

    def test_login_wrong_password_and_lockout(self):
        async def scenario():
            server = await start_server()
            try:
                c = await RawClient(server.port).connect()
                for _ in range(5):
                    r = await c.login("adm", "wrong-password")
                    self.assertEqual(r["code"], P.ERR_VALIDATION)
                r = await c.login("adm", PASSWORD)  # correct password, but locked out
                self.assertEqual(r["code"], P.ERR_RATE_LIMITED)
            finally:
                await server.stop()
        run(scenario())

    def test_protocol_version_mismatch_rejected(self):
        async def scenario():
            server = await start_server()
            try:
                c = await RawClient(server.port).connect()
                r = await c.req("LOGIN", {"user_id": "adm", "password": PASSWORD, "protocol": 1})
                self.assertFalse(r["ok"])
                self.assertEqual(r["code"], P.ERR_FORBIDDEN)
            finally:
                await server.stop()
        run(scenario())

    def test_must_change_password_enforced_server_side(self):
        async def scenario():
            server = await start_server()
            try:
                c = await RawClient(server.port).connect()
                r = await c.login("h_new", "h_new")  # initial password = user ID
                self.assertTrue(r["ok"])
                self.assertEqual(r["user"]["must_change_password"], 1)
                r = await c.req("GET_SNAPSHOT", {})
                self.assertEqual(r["code"], P.ERR_PASSWORD_CHANGE_REQUIRED)
                r = await c.req("SUBMIT_LOG", log_payload(task_id="T-STR"))
                self.assertEqual(r["code"], P.ERR_PASSWORD_CHANGE_REQUIRED)
                # Policy checks
                r = await c.req("CHANGE_PASSWORD", {"new_password": "h_new"})
                self.assertEqual(r["code"], P.ERR_VALIDATION)
                r = await c.req("CHANGE_PASSWORD", {"new_password": "short"})
                self.assertEqual(r["code"], P.ERR_VALIDATION)
                r = await c.req("CHANGE_PASSWORD", {"new_password": "NewPassw0rd"})
                self.assertTrue(r["ok"], r)
                r = await c.req("GET_SNAPSHOT", {})
                self.assertTrue(r["ok"])
                # After the forced change, a further change needs the old password
                r = await c.req("CHANGE_PASSWORD", {"new_password": "Another0ne"})
                self.assertEqual(r["code"], P.ERR_VALIDATION)
                r = await c.req("CHANGE_PASSWORD", {"new_password": "Another0ne", "old_password": "NewPassw0rd"})
                self.assertTrue(r["ok"])
            finally:
                await server.stop()
        run(scenario())

    def test_session_resume_and_revocation(self):
        async def scenario():
            server = await start_server()
            try:
                login = await (await RawClient(server.port).connect()).login("h_urb")
                token = login["session_token"]
                # Resume on a fresh socket works
                c2 = await RawClient(server.port).connect()
                r = await c2.req("RESUME", {"session_token": token, "protocol": P.PROTOCOL_VERSION})
                self.assertTrue(r["ok"])
                self.assertEqual(r["user"]["user_id"], "h_urb")
                # Bogus token fails
                c3 = await RawClient(server.port).connect()
                r = await c3.req("RESUME", {"session_token": "nope", "protocol": P.PROTOCOL_VERSION})
                self.assertEqual(r["code"], P.ERR_SESSION_EXPIRED)
                # Admin resets the Head's password: live sockets are closed, tokens invalidated
                adm = await raw_login(server.port, "adm")
                r = await adm.req("RESET_PASSWORD", {"user_id": "h_urb"})
                self.assertTrue(r["ok"])
                await asyncio.sleep(0.2)
                self.assertEqual(c2.ws.close_code, P.CLOSE_SESSION_REVOKED)
                c4 = await RawClient(server.port).connect()
                r = await c4.req("RESUME", {"session_token": token, "protocol": P.PROTOCOL_VERSION})
                self.assertEqual(r["code"], P.ERR_SESSION_EXPIRED)
            finally:
                await server.stop()
        run(scenario())


class TestIdentityAndScope(unittest.TestCase):
    def test_log_identity_comes_from_session(self):
        async def scenario():
            server = await start_server()
            try:
                h = await raw_login(server.port, "h_urb")
                p = log_payload(employee_id="h_urb2")  # spoof attempt
                r = await h.req("SUBMIT_LOG", p)
                self.assertTrue(r["ok"], r)
                rows = [l for l in server.db.get_progress_logs() if l["log_id"] == p["log_id"]]
                self.assertEqual(rows[0]["employee_id"], "h_urb")
            finally:
                await server.stop()
        run(scenario())

    def test_head_cannot_log_or_request_outside_discipline(self):
        async def scenario():
            server = await start_server()
            try:
                h = await raw_login(server.port, "h_urb")
                r = await h.req("SUBMIT_LOG", log_payload(task_id="T-STR"))
                self.assertEqual(r["code"], P.ERR_FORBIDDEN)
                r = await h.req("REQUEST_TASK_DATE_CHANGE", {"request_id": "DCR-1", "task_id": "T-STR",
                                                             "new_start_date": "2026-10-01", "new_end_date": "2026-10-05",
                                                             "reason": "x"})
                self.assertEqual(r["code"], P.ERR_FORBIDDEN)
            finally:
                await server.stop()
        run(scenario())

    def test_head_sees_only_own_discipline(self):
        async def scenario():
            server = await start_server()
            try:
                h = await raw_login(server.port, "h_str")
                r = await h.req("GET_TASKS", {"discipline_id": "URB"})
                self.assertEqual({t["task_id"] for t in r["tasks"]}, {"T-STR"})
                snap = await h.req("GET_SNAPSHOT", {})
                self.assertEqual({t["task_id"] for t in snap["tasks"]}, {"T-STR"})
                tl = await h.req("GET_PROJECT_TIMELINE", {"project_id": "P1"})
                self.assertEqual({t["task_id"] for t in tl["data"]["tasks"]}, {"T-STR"})
                r = await h.req("GET_TASK_PROPOSALS", {"proposed_by_id": "h_urb"})
                self.assertTrue(all(p["proposed_by_id"] == "h_str" for p in r["proposals"]))
            finally:
                await server.stop()
        run(scenario())

    def test_alert_response_only_by_target(self):
        async def scenario():
            server = await start_server()
            try:
                coo = await raw_login(server.port, "coo1")
                target = await raw_login(server.port, "h_urb")
                other = await raw_login(server.port, "h_str")
                r = await coo.req("COO_NOTIFY", {"target_employee_id": "h_urb", "message": "urgent"})
                self.assertTrue(r["ok"])
                self.assertTrue(r["delivered"])
                nid = r["notification_id"]
                alert = await target.wait_for(lambda d: d.get("type") == "COO_PRIORITY_ALERT")
                self.assertEqual(alert["notification"]["notification_id"], nid)
                r = await other.req("RESPOND_NOTIFICATION", {"notification_id": nid, "response_status": "DENIED"})
                self.assertEqual(r["code"], P.ERR_FORBIDDEN)
                r = await target.req("RESPOND_NOTIFICATION", {"notification_id": nid, "response_status": "ACCEPTED"})
                self.assertTrue(r["ok"])
                ev = await coo.wait_for(lambda d: d.get("type") == "NOTIFICATION_RESPONSE_EVENT")
                self.assertEqual(ev["notification"]["status"], "ACCEPTED")
                # Idempotent resend of the same answer is accepted; a different answer is not
                r = await target.req("RESPOND_NOTIFICATION", {"notification_id": nid, "response_status": "ACCEPTED"})
                self.assertTrue(r["ok"])
                r = await target.req("RESPOND_NOTIFICATION", {"notification_id": nid, "response_status": "DENIED"})
                self.assertEqual(r["code"], P.ERR_VALIDATION)
                # COO_NOTIFY with a client ID is idempotent
                r1 = await coo.req("COO_NOTIFY", {"notification_id": "N-FIXED", "target_employee_id": "h_urb", "message": "m"})
                r2 = await coo.req("COO_NOTIFY", {"notification_id": "N-FIXED", "target_employee_id": "h_urb", "message": "m"})
                self.assertEqual(r1["notification_id"], r2["notification_id"])
                self.assertEqual(len([n for n in server.db.get_priority_notifications() if n["notification_id"] == "N-FIXED"]), 1)
            finally:
                await server.stop()
        run(scenario())

    def test_admin_limited_to_head_accounts(self):
        async def scenario():
            server = await start_server()
            try:
                adm = await raw_login(server.port, "adm")
                r = await adm.req("CREATE_USER", {"user_id": "a2", "full_name": "x", "role": "ADMIN"})
                self.assertEqual(r["code"], P.ERR_FORBIDDEN)
                r = await adm.req("RESET_PASSWORD", {"user_id": "coo1"})
                self.assertEqual(r["code"], P.ERR_FORBIDDEN)
                r = await adm.req("UPDATE_USER", {"user_id": "h_urb", "role": "ADMIN"})
                self.assertEqual(r["code"], P.ERR_FORBIDDEN)
                r = await adm.req("GET_USERS", {})
                self.assertTrue(all(u["role"] == ROLE_HEAD for u in r["users"]))
                r = await adm.req("CREATE_USER", {"user_id": "h9", "full_name": "Head Nine", "role": "HEAD", "discipline_id": "AI"})
                self.assertTrue(r["ok"], r)
                sa = await raw_login(server.port, "sa")
                r = await sa.req("CREATE_USER", {"user_id": "coo2", "full_name": "COO Two", "role": "COO"})
                self.assertTrue(r["ok"], r)
            finally:
                await server.stop()
        run(scenario())


class TestSettingsAndEvents(unittest.TestCase):
    def test_system_settings(self):
        async def scenario():
            server = await start_server()
            try:
                sa = await raw_login(server.port, "sa")
                r = await sa.req("GET_SETTINGS", {})
                self.assertEqual(r["settings"]["backup_retention_count"], 30)
                self.assertIn("port", r["info"])
                r = await sa.req("UPDATE_SETTINGS", {"settings": {"password_min_length": 3}})
                self.assertEqual(r["code"], P.ERR_VALIDATION)
                r = await sa.req("UPDATE_SETTINGS", {"settings": {"backup_time": "25:99"}})
                self.assertEqual(r["code"], P.ERR_VALIDATION)
                r = await sa.req("UPDATE_SETTINGS", {"settings": {"password_min_length": 12, "backup_time": "3:05"}})
                self.assertTrue(r["ok"], r)
                self.assertEqual(r["settings"]["backup_time"], "03:05")
                # New policy applies immediately and is announced at login
                c = await RawClient(server.port).connect()
                login = await c.login("h_new", "h_new")
                self.assertEqual(login["policy"]["password_min_length"], 12)
                r = await c.req("CHANGE_PASSWORD", {"new_password": "Only10Char"})
                self.assertEqual(r["code"], P.ERR_VALIDATION)
                r = await c.req("CHANGE_PASSWORD", {"new_password": "Twelve-Chars1"})
                self.assertTrue(r["ok"], r)
            finally:
                await server.stop()
        run(scenario())

    def test_change_events_are_targeted(self):
        async def scenario():
            server = await start_server()
            try:
                adm = await raw_login(server.port, "adm")
                h_urb = await raw_login(server.port, "h_urb")
                h_str = await raw_login(server.port, "h_str")
                # A new Head account: only oversight roles hear about it
                r = await adm.req("CREATE_USER", {"user_id": "h10", "full_name": "H", "role": "HEAD", "discipline_id": "AI"})
                self.assertTrue(r["ok"])
                with self.assertRaises(TimeoutError):
                    await h_urb.wait_for(lambda d: d.get("type") == "DATA_CHANGED", timeout=0.5)
                # A proposal decision reaches the proposer, not the other Heads
                prop = await h_urb.req("PROPOSE_TASK", {"proposal_id": "PR-1", "project_id": "P1", "task_name": "x"})
                self.assertTrue(prop["ok"])
                await adm.req("REJECT_TASK_PROPOSAL", {"proposal_id": "PR-1", "reason": "no"})
                ev = await h_urb.wait_for(lambda d: d.get("type") == "DATA_CHANGED" and "proposals" in d["scopes"])
                self.assertTrue(ev)
                with self.assertRaises(TimeoutError):
                    await h_str.wait_for(lambda d: d.get("type") == "DATA_CHANGED" and "proposals" in d["scopes"], timeout=0.5)
                # Task changes reach every Head
                await adm.req("ADD_TASK", {"project_id": "P1", "name": "New", "discipline_id": "STR"})
                ev = await h_str.wait_for(lambda d: d.get("type") == "DATA_CHANGED" and "tasks" in d["scopes"])
                self.assertTrue(ev)
            finally:
                await server.stop()
        run(scenario())

    def test_last_super_admin_cannot_be_demoted(self):
        async def scenario():
            server = await start_server()
            try:
                sa = await raw_login(server.port, "sa")
                # 'superadmin' (bootstrap) and 'sa' exist; demote one, the last one is protected
                r = await sa.req("UPDATE_USER", {"user_id": "superadmin", "role": "ADMIN"})
                self.assertTrue(r["ok"], r)
                r = await sa.req("UPDATE_USER", {"user_id": "sa", "role": "COO"})
                self.assertEqual(r["code"], P.ERR_VALIDATION)
            finally:
                await server.stop()
        run(scenario())


class TestSoftDeleteProtocol(unittest.TestCase):
    def test_deleted_project_disappears_for_heads_and_keeps_logs(self):
        async def scenario():
            server = await start_server()
            try:
                h = await raw_login(server.port, "h_urb")
                p = log_payload()
                self.assertTrue((await h.req("SUBMIT_LOG", p))["ok"])
                adm = await raw_login(server.port, "adm")
                r = await adm.req("DELETE_PROJECT", {"project_id": "P1"})
                self.assertTrue(r["ok"], r)
                ev = await h.wait_for(lambda d: d.get("type") == "DATA_CHANGED")
                self.assertIn("projects", ev["scopes"])
                snap = await h.req("GET_SNAPSHOT", {})
                self.assertNotIn("P1", [x["project_id"] for x in snap["projects"]])
                self.assertFalse(any(t["project_id"] == "P1" for t in snap["tasks"]))
                r = await h.req("SUBMIT_LOG", log_payload())
                self.assertEqual(r["code"], P.ERR_VALIDATION)
                r = await adm.req("GET_PROJECTS", {"include_hidden": True})
                self.assertNotIn("P1", [x["project_id"] for x in r["projects"]])
                self.assertTrue(any(l["log_id"] == p["log_id"] for l in server.db.get_progress_logs()))
                r = await adm.req("DELETE_PROJECT", {"project_id": "P1"})
                self.assertEqual(r["code"], P.ERR_VALIDATION)
            finally:
                await server.stop()
        run(scenario())


class TestRobustness(unittest.TestCase):
    def test_malformed_input_never_drops_connection(self):
        async def scenario():
            server = await start_server()
            try:
                c = await RawClient(server.port).connect()
                await c.send_raw("not json")
                r = await c.wait_for(lambda d: d.get("type") == "ERROR")
                self.assertEqual(r["code"], P.ERR_BAD_REQUEST)
                r = await c.req("LOGIN", {"user_id": None})
                self.assertFalse(r["ok"])
                await c.send_raw(json.dumps({"action": "LOGIN", "payload": "string", "msg_id": "m1"}))
                r = await c.wait_for(lambda d: d.get("msg_id") == "m1")
                self.assertEqual(r["code"], P.ERR_BAD_REQUEST)
                h = await raw_login(server.port, "h_urb")
                r = await h.req("SUBMIT_LOG", log_payload(progress_stage="Almost"))
                self.assertFalse(r["ok"])
                r = await h.req("SUBMIT_LOG", log_payload(task_id=None))
                self.assertEqual(r["code"], P.ERR_VALIDATION)
                # Connection still usable
                self.assertTrue((await h.req("GET_SNAPSHOT", {}))["ok"])
                self.assertTrue((await c.req("PING", {}))["ok"])
            finally:
                await server.stop()
        run(scenario())

    def test_write_errors_are_reported(self):
        async def scenario():
            server = await start_server()
            try:
                adm = await raw_login(server.port, "adm")
                r = await adm.req("ADD_PROJECT", {"project_id": "P1", "name": "dup"})
                self.assertFalse(r["ok"])
                self.assertIn("already exists", r["error"])
                r = await adm.req("REJECT_TASK_PROPOSAL", {"proposal_id": "X"})
                self.assertEqual(r["code"], P.ERR_VALIDATION)
            finally:
                await server.stop()
        run(scenario())


class TestExportStreaming(unittest.TestCase):
    def test_export_is_chunked_and_does_not_block_other_clients(self):
        async def scenario():
            server = await start_server(export_chunk_bytes=16 * 1024)
            try:
                con = sqlite3.connect(server.db.db_path)
                con.executemany(
                    """INSERT INTO progress_logs (log_id, employee_id, project_id, task_id, start_date, end_date,
                       progress_stage, notes, assigned_to, discipline_id) VALUES (?,?,?,?,?,?,?,?,?,?)""",
                    [(f"L{i}", "h_urb", "P1", "T-URB", "2026-10-01", "2026-10-01", "In Progress",
                      f"note {i} " * 10, "Eng. X", "URB") for i in range(3000)])
                con.commit()
                con.close()

                exporter = await raw_login(server.port, "coo1")
                bystander = await raw_login(server.port, "h_urb")
                msg_id = "EXP-1"
                await exporter.ws.send(json.dumps({"action": "EXPORT_EXCEL_REQUEST", "payload": {}, "msg_id": msg_id}))

                # While the workbook is being built, another client keeps getting fast answers.
                latencies = []
                loop = asyncio.get_running_loop()
                for _ in range(5):
                    t = loop.time()
                    r = await bystander.req("GET_SNAPSHOT", {}, timeout=5)
                    self.assertTrue(r["ok"])
                    latencies.append(loop.time() - t)
                self.assertLess(max(latencies), 2.0)

                chunks = []
                final = None
                while final is None:
                    d = json.loads(await asyncio.wait_for(exporter.ws.recv(), timeout=120))
                    if d.get("type") == "EXPORT_EXCEL_CHUNK" and d["msg_id"] == msg_id:
                        chunks.append(d)
                    elif d.get("msg_id") == msg_id:
                        final = d
                self.assertTrue(final["ok"], final)
                self.assertGreater(final["total_chunks"], 1)
                self.assertEqual(len(chunks), final["total_chunks"])
                import base64, io
                from openpyxl import load_workbook
                raw = b"".join(base64.b64decode(c["data_b64"]) for c in sorted(chunks, key=lambda c: c["index"]))
                self.assertEqual(len(raw), final["size"])
                wb = load_workbook(io.BytesIO(raw), read_only=True)
                self.assertIn("All Logs", wb.sheetnames)
            finally:
                await server.stop()
        run(scenario())


if __name__ == "__main__":
    unittest.main()

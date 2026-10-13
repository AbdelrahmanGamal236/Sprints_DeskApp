"""M4: Excel export content, filters, analytics numbers and project timeline."""
import io
import shutil
import sqlite3
import tempfile
import time
import unittest
import uuid
from pathlib import Path

from openpyxl import load_workbook

from app.core.config import ROLE_HEAD, ROLE_SUPER_ADMIN
from app.database.db_manager import DBManager
from app.services.excel_service import ExcelExportService, HEADERS, TASK_HEADERS


class ReportBase(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="deskapp_reports_"))
        db = self.db = DBManager(central_db_path=self.tmp / "c.db")
        db.add_discipline("MEP", "Mech/Elec: Services")  # name with characters Excel forbids in sheet titles
        db.create_user(ROLE_SUPER_ADMIN, "hu", "Urban Head", ROLE_HEAD, "URB")
        db.create_user(ROLE_SUPER_ADMIN, "hs", "Structure Head", ROLE_HEAD, "STR")
        db.create_user(ROLE_SUPER_ADMIN, "hm", "MEP Head", ROLE_HEAD, "MEP")
        db.add_project("P1", "Tower")
        db.add_project("P2", "Mall")
        db.add_task("P1", "U1", "Masterplan", "URB")
        db.add_task("P1", "U2", "Landscape", "URB")
        db.add_task("P1", "S1", "Foundations", "STR")
        db.add_task("P1", "S2", "Columns", "STR")
        db.add_task("P2", "M1", "HVAC", "MEP")
        self.svc = ExcelExportService(db)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def log(self, emp, task, project="P1", submitted=None, **extra):
        p = {"log_id": f"LOG-{uuid.uuid4().hex}", "employee_id": emp, "project_id": project, "task_id": task,
             "progress_stage": "In Progress"}
        p.update(extra)
        saved = self.db.submit_progress_log(p)
        if submitted:  # simulate a log submitted on another day
            con = sqlite3.connect(self.db.db_path)
            con.execute("UPDATE progress_logs SET submitted_at = ? WHERE log_id = ?", (submitted, saved["log_id"]))
            con.commit()
            con.close()
        return saved

    def workbook(self, **filters):
        return load_workbook(io.BytesIO(self.svc.export_logs_to_excel(**filters)))

    @staticmethod
    def rows(ws):
        values = list(ws.iter_rows(values_only=True))
        header = list(values[0])
        return [dict(zip(header, v)) for v in values[1:]]


class TestExcelContent(ReportBase):
    def test_sheets_headers_and_labels(self):
        self.db.update_task("U1", planned_start_date="2026-10-01", planned_end_date="2026-10-14")
        self.log("hu", "U1", submitted="2026-10-10 08:00:00")  # first entry of each task: New Task
        self.log("hu", "U2", submitted="2026-10-10 08:05:00")
        self.log("hu", "U1", entry_type="REVISION", assigned_to="Eng. Sara", submitted="2026-10-13 09:45:12")
        self.log("hu", "U1", entry_type="REVISION")
        self.log("hu", "U2", entry_type="CONTINUATION")
        self.log("hs", "S1", progress_stage="Done")
        self.log("hm", "M1", project="P2")
        self.db.toggle_hide_task("U2", 0)
        self.db.delete_task("S1", "adm")
        self.db.toggle_hide_discipline("MEP", 0)

        wb = self.workbook(generated_by="Admin (adm)")
        self.assertEqual(wb.sheetnames[0], "All Logs")
        self.assertEqual(wb.sheetnames[-1], "Report Info")
        self.assertIn("Mech-Elec- Services", wb.sheetnames)  # sanitized sheet title
        self.assertIn("Urban", wb.sheetnames)
        all_logs = self.rows(wb["All Logs"])
        self.assertEqual([c.value for c in wb["All Logs"][1]], HEADERS)
        self.assertEqual(len(all_logs), 7)

        by_task = {}
        for r in all_logs:
            by_task.setdefault(r["Task ID"], []).append(r)
        self.assertEqual(sorted(r["Revision Tag"] or "" for r in by_task["U1"]), ["", "R1", "R2"])
        self.assertEqual(sorted(r["Entry Type"] for r in by_task["U1"]), ["New", "Revision", "Revision"])
        self.assertNotIn("Progress %", HEADERS)
        sara = next(r for r in by_task["U1"] if r["Assigned To"] == "Eng. Sara")
        self.assertEqual((sara["Submitted Date"], sara["Submitted Time"]), ("2026-10-13", "09:45:12"))
        self.assertEqual((sara["Task Start Date"], sara["Task End Date"]), ("2026-10-01", "2026-10-14"))
        self.assertEqual(sara["Status"], "In Progress")
        self.assertEqual(by_task["S1"][0]["Status"], "Done")
        self.assertIn("Eng. Sara", [r["Assigned To"] for r in by_task["U1"]])
        self.assertEqual(by_task["U2"][0]["Task"], "Landscape (Hidden)")
        self.assertEqual(by_task["U2"][0]["Entry Type"], "Continuation")
        self.assertEqual(by_task["S1"][0]["Task"], "Foundations (Deleted)")
        self.assertEqual(by_task["M1"][0]["Discipline"], "Mech/Elec: Services (Hidden)")
        self.assertEqual(by_task["U1"][0]["Head Discipline"], "Urban")
        self.assertEqual(by_task["U1"][0]["Head Name"], "Urban Head")

        # Per-discipline sheets contain only their discipline's rows
        self.assertEqual({r["Task ID"] for r in self.rows(wb["Urban"])}, {"U1", "U2"})
        info = {r[0]: r[1] for r in wb["Report Info"].iter_rows(values_only=True) if r and r[0]}
        self.assertEqual(info["Rows exported"], 7)
        self.assertEqual(info["Generated by"], "Admin (adm)")

    def test_filters_and_submission_dates(self):
        self.log("hu", "U1", submitted="2026-10-05 17:30:00")
        self.log("hu", "U2", submitted="2026-10-08 08:00:00")
        self.log("hs", "S1", submitted="2026-09-25 23:59:59")
        self.log("hm", "M1", project="P2", submitted="2026-10-03 12:00:00")

        def ids(**f):
            return sorted(r["Task ID"] for r in self.rows(self.workbook(**f)["All Logs"]))

        self.assertEqual(ids(project_id="P2"), ["M1"])
        self.assertEqual(ids(discipline_id="URB"), ["U1", "U2"])
        self.assertEqual(ids(employee_id="hs"), ["S1"])
        # By submission day, both ends inclusive (a log at 17:30 on the To date is included)
        self.assertEqual(ids(start_date="2026-10-04", end_date="2026-10-08"), ["U1", "U2"])
        self.assertEqual(ids(start_date="2026-10-05", end_date="2026-10-05"), ["U1"])
        self.assertEqual(ids(start_date="2026-10-06", end_date="2026-10-10"), ["U2"])
        self.assertEqual(ids(end_date="2026-09-25"), ["S1"])
        self.assertEqual(ids(discipline_id="URB", start_date="2026-10-06"), ["U2"])
        with self.assertRaises(ValueError):
            self.svc.export_logs_to_excel(start_date="10/01/2026")
        with self.assertRaises(ValueError):
            self.svc.export_logs_to_excel(start_date="2026-10-10", end_date="2026-10-01")

        wb = self.workbook(discipline_id="URB")
        self.assertEqual(wb.sheetnames, ["All Logs", "Task Summary", "Report Info"])

    def test_task_summary(self):
        """Every task in the filters' scope is listed, with its dates and latest status, even without logs."""
        self.db.update_task("U1", planned_start_date="2026-10-01", planned_end_date="2026-10-14")
        self.log("hu", "U1", submitted="2026-10-12 09:00:00")
        self.log("hu", "U1", progress_stage="Done", submitted="2026-10-13 15:30:00")
        self.db.toggle_hide_task("U2", 0)

        def summary(**f):
            ws = self.workbook(**f)["Task Summary"]
            self.assertEqual([c.value for c in ws[1]], TASK_HEADERS)
            return {r["Task ID"]: r for r in self.rows(ws)}

        urb = summary(discipline_id="URB")
        self.assertEqual(set(urb), {"U1", "U2"})
        self.assertEqual((urb["U1"]["Task Start Date"], urb["U1"]["Task End Date"]), ("2026-10-01", "2026-10-14"))
        self.assertEqual((urb["U1"]["Latest Status"], urb["U1"]["Last Update"]), ("Done", "2026-10-13 15:30:00"))
        self.assertEqual((urb["U1"]["Last Updated By"], urb["U1"]["Logs in Report"]), ("Urban Head", 2))
        self.assertEqual((urb["U2"]["Latest Status"], urb["U2"]["Logs in Report"]), ("No update yet", 0))
        self.assertEqual(urb["U2"]["Task"], "Landscape (Hidden)")
        # Date range limits the logs the status is taken from, not the task list
        early = summary(discipline_id="URB", end_date="2026-10-12")
        self.assertEqual((early["U1"]["Latest Status"], early["U1"]["Logs in Report"]), ("In Progress", 1))
        self.assertEqual(set(summary(project_id="P2")), {"M1"})
        self.assertEqual(set(summary(employee_id="hs")), {"S1", "S2"})  # a Head's own discipline tasks
        # Counts reported to the client (used for the "Nothing to Export" message)
        _data, counts = self.svc.build_report(project_id="P2", discipline_id="URB")
        self.assertEqual(counts, {"logs": 0, "tasks": 0})
        _data, counts = self.svc.build_report(discipline_id="URB")
        self.assertEqual(counts, {"logs": 2, "tasks": 2})

    def test_empty_export_still_valid(self):
        wb = self.workbook(project_id="P2")
        self.assertEqual(self.rows(wb["All Logs"]), [])


class TestAnalytics(ReportBase):
    def test_counts_are_not_multiplied_by_tasks(self):
        for _ in range(3):
            self.log("hu", "U1")
        self.log("hu", "U1", entry_type="REVISION")
        self.log("hu", "U2", progress_stage="Done")
        self.log("hs", "S1", progress_stage="Hold")
        data = self.db.get_analytics_data()
        p1 = next(p for p in data["project_breakdown"] if p["project_id"] == "P1")
        self.assertEqual(p1["total_tasks"], 4)
        self.assertEqual(p1["total_revisions"], 1)
        # status by each task's latest log: U1 In Progress, U2 Done, S1 Hold, S2 not started
        self.assertEqual((p1["in_prog_count"], p1["done_count"], p1["hold_count"], p1["not_started_count"]), (1, 1, 1, 1))
        self.assertEqual(data["total_logs"], 6)
        hu = next(e for e in data["employee_breakdown"] if e["user_id"] == "hu")
        # One New Task entry per task (U1, U2); the other U1 entries became Continuations
        self.assertEqual((hu["total_logs"], hu["new_logs"], hu["rev_logs"]), (5, 2, 1))
        hm = next(e for e in data["employee_breakdown"] if e["user_id"] == "hm")
        self.assertEqual(hm["total_logs"], 0)

    def test_deleted_items_excluded_from_live_counts(self):
        self.db.delete_task("S2", "adm")
        self.db.delete_project("P2", "adm")
        data = self.db.get_analytics_data()
        self.assertEqual(data["active_projects"], 1)
        self.assertEqual(next(p for p in data["project_breakdown"] if p["project_id"] == "P1")["total_tasks"], 3)


class TestTimeline(ReportBase):
    def test_latest_log_per_task(self):
        self.log("hu", "U1", progress_stage="Done")
        self.log("hu", "U1", entry_type="REVISION", assigned_to="Eng. Late")  # reopened: back to In Progress
        self.log("hu", "U2", progress_stage="Done")
        tl = self.db.get_project_timeline_data("P1")
        tasks = {t["task_id"]: t for t in tl["tasks"]}
        self.assertEqual(tasks["U1"]["progress_percentage"], 0)  # progress follows the latest stage only
        self.assertEqual(tasks["U1"]["revision_count"], 1)
        self.assertEqual(tasks["U1"]["latest_assigned_to"], "Eng. Late")
        self.assertEqual(tasks["U2"]["progress_percentage"], 100)
        self.assertEqual(tasks["S1"]["progress_stage"], "Not Started")
        self.assertEqual(tl["completed_tasks"], 1)
        self.assertEqual(tl["overall_progress_percentage"], round(100 / 4))


class TestExportPerformance(ReportBase):
    def test_20k_logs_export_time(self):
        con = sqlite3.connect(self.db.db_path)
        con.executemany(
            """INSERT INTO progress_logs (log_id, employee_id, project_id, task_id, start_date, end_date,
               progress_stage, notes, assigned_to, discipline_id, submitted_at) VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
            [(f"L{i}", "hu", "P1", "U1", "2026-10-01", "2026-10-01", "In Progress", "note text " * 15, "Eng. X",
              "URB", "2026-10-01 10:00:00") for i in range(20000)])
        con.commit()
        con.close()
        t = time.time()
        data = self.svc.export_logs_to_excel()
        elapsed = time.time() - t
        print(f"\n20k-log export: {elapsed:.1f}s, {len(data) // 1024} KB")
        self.assertLess(elapsed, 60)  # was ~65s for 20k rows before M4 (50s for 15k)


if __name__ == "__main__":
    unittest.main()

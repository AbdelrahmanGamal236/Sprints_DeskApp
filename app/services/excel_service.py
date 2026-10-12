import io
import re
from datetime import datetime

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.database.db_manager import DBManager, parse_date, check_date_order

# One row per submitted log: when it was submitted, by whom, on which task (with the task's planned
# period at that time) and the status reported. Newest first.
HEADERS = [
    "Submitted Date", "Submitted Time", "Head ID", "Head Name", "Head Discipline",
    "Project", "Task ID", "Task", "Discipline", "Task Start Date", "Task End Date",
    "Status", "Entry Type", "Revision Tag", "Assigned To", "Notes", "Log ID",
]
ENTRY_LABELS = {"NEW": "New", "CONTINUATION": "Continuation", "REVISION": "Revision"}
MAX_WIDTH = 60

_HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
_BORDER = Border(*(Side(style="thin", color="D1D5DB"),) * 4)
_MUTED_FONT = Font(name="Calibri", size=10, italic=True, color="6B7280")
_CENTER = Alignment(horizontal="center", vertical="center")


def _label(name, ident, hidden=False, deleted=False, fallback="Unknown"):
    text = name or ident or fallback
    if deleted:
        text += " (Deleted)"
    elif hidden:
        text += " (Hidden)"
    return text


def _sheet_title(name: str, used: set) -> str:
    """Excel sheet names: max 31 chars, no []:*?/\\ and unique (case-insensitive)."""
    base = re.sub(r"[\[\]:*?/\\]", "-", name or "Sheet").strip() or "Sheet"
    base = base[:31]
    title, n = base, 2
    while title.lower() in used:
        suffix = f" ({n})"
        title = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(title.lower())
    return title


class ExcelExportService:
    def __init__(self, db_manager: DBManager):
        self.db = db_manager

    @staticmethod
    def build_rows(logs: list[dict]) -> list[tuple[list, bool]]:
        """Turn log records into sheet rows; second item marks rows of hidden/deleted items (shown muted)."""
        rows = []
        for r in logs:
            entry = (r.get("entry_type") or "NEW").upper()
            rev = r.get("revision_number") or 0
            inactive = bool(r.get("project_hidden") or r.get("project_deleted") or
                            r.get("task_hidden") or r.get("task_deleted"))
            submitted = r.get("submitted_at") or ""
            rows.append(([
                submitted[:10],
                submitted[11:19],
                r["employee_id"],
                r.get("employee_name") or r["employee_id"],
                r.get("head_discipline_name") or r.get("discipline_id") or "Unassigned",
                _label(r.get("project_name"), r["project_id"], r.get("project_hidden"), r.get("project_deleted"),
                       "Unknown Project"),
                r["task_id"],
                _label(r.get("task_name"), r["task_id"], r.get("task_hidden"), r.get("task_deleted"), "Unknown Task"),
                _label(r.get("task_discipline_name"), r.get("task_discipline_id"), r.get("task_discipline_hidden"),
                       False, "Unassigned"),
                r["start_date"],
                r["end_date"],
                r.get("progress_stage") or "",
                ENTRY_LABELS.get(entry, entry.title()),
                f"R{rev}" if entry == "REVISION" and rev else "",
                r.get("assigned_to") or "",
                r.get("notes") or "",
                r["log_id"],
            ], inactive))
        return rows

    def _write_sheet(self, wb, title: str, rows: list):
        ws = wb.create_sheet(title=title)
        widths = [len(h) for h in HEADERS]
        for values, _ in rows:
            for i, v in enumerate(values):
                widths[i] = max(widths[i], len(str(v)))
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = min(max(w + 3, 10), MAX_WIDTH)
        ws.freeze_panes = "A2"

        header_cells = []
        for h in HEADERS:
            c = WriteOnlyCell(ws, value=h)
            c.font, c.fill, c.alignment, c.border = _HEADER_FONT, _HEADER_FILL, _CENTER, _BORDER
            header_cells.append(c)
        ws.append(header_cells)

        # Plain values for normal rows keep large exports fast; only rows of hidden/deleted
        # items get per-cell styling (grey italic).
        for values, inactive in rows:
            if not inactive:
                ws.append(values)
                continue
            cells = []
            for v in values:
                c = WriteOnlyCell(ws, value=v)
                c.font = _MUTED_FONT
                cells.append(c)
            ws.append(cells)

    def export_logs_to_excel(self, output_path: str = None, project_id: str = None,
                             discipline_id: str = None, employee_id: str = None,
                             start_date: str = None, end_date: str = None,
                             generated_by: str = None) -> bytes:
        """
        Builds the progress-log workbook on the server.
        Sheets: "All Logs", one sheet per discipline (when not filtered to one discipline), and "Report Info".
        Date range: a log is included when it was submitted on a day from start_date to end_date.
        Logs of hidden/deleted projects, tasks and disciplines are included and labelled.
        """
        start_date = parse_date(start_date, "From date", required=False)
        end_date = parse_date(end_date, "To date", required=False)
        check_date_order(start_date, end_date, "From date")

        logs = self.db.get_progress_logs(project_id=project_id, discipline_id=discipline_id,
                                         employee_id=employee_id, start_date=start_date, end_date=end_date)
        rows = self.build_rows(logs)

        wb = openpyxl.Workbook(write_only=True)
        used = set()
        self._write_sheet(wb, _sheet_title("All Logs", used), rows)

        if not discipline_id:
            disc_names = {d["discipline_id"]: d["name"] for d in self.db.get_disciplines(include_hidden=True)}
            grouped: dict[str, list] = {}
            for log, row in zip(logs, rows):
                d_id = log.get("task_discipline_id") or log.get("discipline_id") or "Unassigned"
                grouped.setdefault(d_id, []).append(row)
            for d_id in sorted(grouped, key=lambda d: disc_names.get(d, d)):
                self._write_sheet(wb, _sheet_title(disc_names.get(d_id, d_id), used), grouped[d_id])

        self._write_info_sheet(wb, _sheet_title("Report Info", used), {
            "Project": project_id or "All projects",
            "Discipline": discipline_id or "All disciplines",
            "Head": employee_id or "All Heads",
            "From date": start_date or "-",
            "To date": end_date or "-",
            "Date rule": "A log is included when it was submitted between the From and To dates (inclusive)",
            "Rows exported": len(rows),
            "Generated at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "Generated by": generated_by or "-",
        })

        buf = io.BytesIO()
        wb.save(buf)
        excel_bytes = buf.getvalue()
        if output_path:
            with open(output_path, "wb") as f:
                f.write(excel_bytes)
        return excel_bytes

    @staticmethod
    def _write_info_sheet(wb, title: str, info: dict):
        ws = wb.create_sheet(title=title)
        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 70
        head = WriteOnlyCell(ws, value="DeskApp Progress Log Report")
        head.font = Font(name="Calibri", size=13, bold=True)
        ws.append([head])
        ws.append([])
        for k, v in info.items():
            key = WriteOnlyCell(ws, value=k)
            key.font = Font(name="Calibri", size=10, bold=True)
            ws.append([key, v])
        ws.append([])
        note = WriteOnlyCell(ws, value="(Hidden) / (Deleted) mark items an Admin has hidden or deleted; "
                                       "their logs are kept and shown in grey.")
        note.font = _MUTED_FONT
        ws.append([note])
        footer = WriteOnlyCell(ws, value="All Rights Reserved to Eng. Abdelrahman Gamal & Eng. Tolimy")
        footer.font = Font(name="Calibri", size=10, italic=True, bold=True, color="475569")
        ws.append([footer])
